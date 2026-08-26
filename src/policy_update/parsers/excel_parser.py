from __future__ import annotations
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any, ClassVar
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from .base_parser import BaseParser, BlockType, ContentBlock, DocumentRole, ParsedDocument, ParseRequest, ParseWarning, SourceLocation
from .parser_exceptions import CorruptedFile, PasswordProtectedFile
from config import EXCEL_HEADER_MIN_SCORE, EXCEL_HEADER_SEARCH_ROWS, EXCEL_MAX_DATA_ROWS_PER_SHEET, EXCEL_MAX_SCAN_ROWS, EXCEL_TABLE_BLANK_GAP_ROWS

@dataclass
class _OrderCounter:
    """Simple counter utility that increments and returns sequential integer values."""
    
    value: int = 0

    def next(self) -> int:
        current = self.value
        self.value += 1
        return current

def _is_percent_format(number_format: str | None) -> bool:
    """Check if a cell's number format represents percentage."""
    
    return bool(number_format) and "%" in number_format

def _cell_display(value: Any, number_format: str | None = None) -> str:
    """Convert a raw Excel cell value into human-readable text, respecting date, time, boolean, float, and percent formats."""
    
    if value is None:
        return ""

    if isinstance(value, datetime):
        if value.time().hour == 0 and value.time().minute == 0 and value.time().second == 0:
            return value.strftime("%d/%m/%Y")
        return value.strftime("%d/%m/%Y %H:%M")

    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")

    if isinstance(value, bool):
        return str(value)

    if isinstance(value, (int, float)) and _is_percent_format(number_format):
        percent_value = value * 100
        if float(percent_value).is_integer():
            return f"{int(percent_value)}%"
        return f"{percent_value:.2f}".rstrip("0").rstrip(".") + "%"

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        
        return f"{value:.4f}".rstrip("0").rstrip(".")

    return str(value).strip()

def _cell_raw(value: Any) -> Any:
    """Return the raw cell value for computation, keeping full precision and original scale. Dates are converted to ISO strings."""
    
    if isinstance(value, (datetime, date)):
        return value.isoformat()

    return value

def _build_merged_value_map(ws: Worksheet) -> dict[tuple[int, int], Any]:
    """Map all cells in a merged range to the top-left cell's value."""

    merged_map: dict[tuple[int, int], Any] = {}

    for merged_range in ws.merged_cells.ranges:
        top_left_value = ws.cell(row=merged_range.min_row, column=merged_range.min_col).value

        if top_left_value is None:
            continue

        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                merged_map[(row, col)] = top_left_value

    return merged_map

def _find_used_extent(ws: Worksheet, max_scan_rows: int) -> tuple[int, int]:
    """Find the last row and column with data, scanning up to max_scan_rows for efficiency."""

    last_row = 0
    last_col = 0

    scan_limit = min(ws.max_row or 0, max_scan_rows)

    if scan_limit <= 0:
        return 0, 0

    for row in ws.iter_rows(min_row=1, max_row=scan_limit):
        for cell in row:
            if cell.value is not None and str(cell.value).strip() != "":
                if cell.row > last_row:
                    last_row = cell.row
                if cell.column > last_col:
                    last_col = cell.column

    return last_row, last_col

def _score_header_row(cells: list[Any]) -> float:
    """Compute a heuristic score (0–1) for how header-like a row is based on fill, text, and distinct value ratios."""

    non_empty = [c for c in cells if c is not None and str(c).strip() != ""]

    if len(non_empty) < 2:
        return 0.0

    text_like = [c for c in non_empty if isinstance(c, str) and not str(c).strip().replace(".", "").isdigit()]
    distinct = len({str(c).strip() for c in non_empty})

    fill_ratio = len(non_empty) / len(cells) if cells else 0.0
    text_ratio = len(text_like) / len(non_empty)
    distinct_ratio = distinct / len(non_empty)

    return fill_ratio * 0.3 + text_ratio * 0.5 + distinct_ratio * 0.2

class ExcelParser(BaseParser):
    parser_name: ClassVar[str] = "excel_parser"
    parser_version: ClassVar[str] = "1.0.0"
    supported_extensions: ClassVar[frozenset[str]] = frozenset({".xlsx", ".xlsm"})
    supported_mime_types: ClassVar[frozenset[str]] = frozenset(
        {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel.sheet.macroEnabled.12",
        }
    )
    supported_roles: ClassVar[frozenset[DocumentRole]] = frozenset(DocumentRole)

    def __init__(self, *, sheet_overrides: dict[str, dict[str, Any]] | None = None, include_hidden_sheets: bool = True,) -> None:
        """Initialize ExcelParser with optional sheet overrides (custom header/skip rows per sheet) and a flag to include hidden sheets in parsing."""

        self.sheet_overrides = sheet_overrides or {}
        self.include_hidden_sheets = include_hidden_sheets

    def _parse(self, request: ParseRequest) -> ParsedDocument:
        workbook = self._open_workbook(request)

        blocks: list[ContentBlock] = []
        warnings: list[ParseWarning] = []
        order = _OrderCounter()
        table_counter = _OrderCounter()
        sheet_count = 0

        try:
            for worksheet in workbook.worksheets:
                is_hidden = worksheet.sheet_state != "visible"

                if is_hidden and not self.include_hidden_sheets:
                    warnings.append(
                        ParseWarning(code="SHEET_HIDDEN_SKIPPED",
                                     message=f"Sheet '{worksheet.title}' is hidden and was skipped.",
                                     location=SourceLocation(sheet=worksheet.title),))
                    continue

                sheet_count += 1

                if is_hidden:
                    warnings.append(
                        ParseWarning(
                            code="SHEET_HIDDEN_INCLUDED",
                            message=(f"Sheet '{worksheet.title}' is hidden but was parsed anyway"),
                            location=SourceLocation(sheet=worksheet.title),))

                sheet_blocks, sheet_warnings = self._process_sheet(
                    worksheet=worksheet,
                    order=order,
                    table_counter=table_counter,
                    is_hidden=is_hidden,
                )

                blocks.extend(sheet_blocks)
                warnings.extend(sheet_warnings)

        except (CorruptedFile, PasswordProtectedFile):
            raise

        except Exception as exc:
            raise CorruptedFile(
                message="Error reading Excel content",
                source_id=request.source_ref.source_id,
                cause=exc,
            ) from exc

        finally:
            workbook.close()

        metadata = ParsedDocument.create_metadata(
            parser_name=self.parser_name,
            parser_version=self.parser_version,
            request=request,
            sheet_count=sheet_count,
        )

        return ParsedDocument(
            source_ref=request.source_ref,
            blocks=blocks,
            warnings=warnings,
            metadata=metadata,
        )

    @staticmethod
    def _open_workbook(request: ParseRequest):
        try:
            return load_workbook(request.file_stream, data_only=True, read_only=False)

        except InvalidFileException as exc:
            raise CorruptedFile(
                message="Not a valid Excel (.xlsx/.xlsm) file",
                source_id=request.source_ref.source_id,
                cause=exc,
            ) from exc

        except Exception as exc:
            message = str(exc).lower()

            if "password" in message or "encrypt" in message:
                raise PasswordProtectedFile(
                    source_id=request.source_ref.source_id,
                    cause=exc,
                ) from exc

            raise CorruptedFile(
                message="Cannot open Excel file",
                source_id=request.source_ref.source_id,
                cause=exc,
            ) from exc

    def _process_sheet(self, *, worksheet: Worksheet, order: _OrderCounter, table_counter: _OrderCounter, is_hidden: bool = False,) -> tuple[list[ContentBlock], list[ParseWarning]]:
        """Parse a single worksheet into content blocks, detecting headers, titles, tables, and warnings."""
        
        blocks: list[ContentBlock] = []
        warnings: list[ParseWarning] = []

        last_row, last_col = _find_used_extent(worksheet, EXCEL_MAX_SCAN_ROWS)

        if last_row == 0 or last_col == 0:
            warnings.append(
                ParseWarning(
                    code="SHEET_EMPTY",
                    message=f"Sheet '{worksheet.title}' has no data.",
                    location=SourceLocation(sheet=worksheet.title),
                )
            )
            return blocks, warnings

        merged_map = _build_merged_value_map(worksheet)

        def get_value(row: int, col: int) -> Any:
            cell_value = worksheet.cell(row=row, column=col).value
            if cell_value is not None:
                return cell_value
            return merged_map.get((row, col))

        def get_format(row: int, col: int) -> str | None:
            return worksheet.cell(row=row, column=col).number_format

        override = self.sheet_overrides.get(worksheet.title)

        scan_start = 1
        is_first_table = True

        while scan_start <= last_row:
            if is_first_table and override and override.get("header_rows"):
                header_rows = sorted(override["header_rows"])
                header_row_idx: int | None = header_rows[0]
                skip_rows: set[int] = set(override.get("skip_rows", []))
                header_score = None
            else:
                header_row_idx, header_span, header_score = self._detect_header(
                    get_value=get_value,
                    first_row=scan_start,
                    last_row=last_row,
                    last_col=last_col,
                )
                header_rows = (
                    list(range(header_row_idx, header_row_idx + header_span))
                    if header_row_idx is not None
                    else []
                )
                skip_rows = set()

            if header_row_idx is None:
                if header_score is not None:
                    warnings.append(
                        ParseWarning(
                            code="HEADER_NOT_DETECTED",
                            message=(
                                f"Sheet '{worksheet.title}' (from row {scan_start}): no confident "
                                f"header row found in the next {EXCEL_HEADER_SEARCH_ROWS} rows "
                                f"(best score {header_score:.2f} < {EXCEL_HEADER_MIN_SCORE}). "
                                f"Remaining rows are emitted without column labels."
                            ),
                            location=SourceLocation(sheet=worksheet.title, row=scan_start),
                        )
                    )

                table_index = table_counter.next()
                blocks.extend(
                    self._emit_freeform_rows(
                        worksheet=worksheet,
                        get_value=get_value,
                        start_row=scan_start,
                        last_row=last_row,
                        last_col=last_col,
                        order=order,
                        table_index=table_index,
                        warnings=warnings,
                        is_hidden=is_hidden,
                    )
                )
                break

            table_index = table_counter.next()
            title_lines: list[str] = []
            title_cells: list[tuple[int, int]] = []  

            for row_idx in range(scan_start, header_row_idx):
                row_values = [get_value(row_idx, c) for c in range(1, last_col + 1)]
                non_empty = [
                    (c, v) for c, v in enumerate(row_values, start=1)
                    if v is not None and str(v).strip() != ""
                ]

                if len(non_empty) == 1:
                    col_idx, value = non_empty[0]
                    title_lines.append(_cell_display(value))
                    title_cells.append((row_idx, col_idx))

            if title_lines:
                title_text = "\n".join(title_lines)
                block_order = order.next()

                title_rows = [r for r, _ in title_cells]
                title_cols = [c for _, c in title_cells]
                min_row, max_row = min(title_rows), max(title_rows)
                min_col, max_col = min(title_cols), max(title_cols)

                blocks.append(
                    ContentBlock(
                        block_id=f"excel-{block_order:06d}",
                        block_type=BlockType.HEADING,
                        raw_text=title_text,
                        normalized_text=title_text,
                        order=block_order,
                        location=SourceLocation(
                            sheet=worksheet.title,
                            table_index=table_index,
                            row=min_row,
                            column=min_col if len(title_cells) == 1 else None,
                            cell_range=f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}",
                        ),
                        metadata={
                            "heading_label": worksheet.title,
                            "heading_level": 1,
                            "sheet_hidden": is_hidden,
                        },
                    )
                )

            headers = self._build_headers(
                get_value=get_value,
                header_rows=header_rows,
                last_col=last_col,
            )

            data_start_row = max(header_rows) + 1

            table_end_row = self._find_table_end(
                get_value=get_value,
                start_row=data_start_row,
                last_row=last_row,
                last_col=last_col,
            )

            blocks.extend(
                self._emit_table_rows(
                    worksheet=worksheet,
                    get_value=get_value,
                    get_format=get_format,
                    headers=headers,
                    start_row=data_start_row,
                    last_row=table_end_row,
                    last_col=last_col,
                    order=order,
                    table_index=table_index,
                    warnings=warnings,
                    skip_rows=skip_rows,
                    is_hidden=is_hidden,
                )
            )

            scan_start = table_end_row + 1
            is_first_table = False

        return blocks, warnings

    @staticmethod
    def _find_table_end(*, get_value, start_row: int, last_row: int, last_col: int,) -> int:
        """Detect the end of a table by finding consecutive blank rows; otherwise run to last_row."""
        
        consecutive_blank = 0
        last_non_blank = start_row - 1

        for row_idx in range(start_row, last_row + 1):
            row_values = [get_value(row_idx, c) for c in range(1, last_col + 1)]
            is_blank = all(v is None or str(v).strip() == "" for v in row_values)

            if is_blank:
                consecutive_blank += 1
                if consecutive_blank >= EXCEL_TABLE_BLANK_GAP_ROWS:
                    return last_non_blank
            else:
                consecutive_blank = 0
                last_non_blank = row_idx

        return last_row

    @staticmethod
    def _detect_header(*, get_value, first_row: int, last_row: int, last_col: int,) -> tuple[int | None, int, float]:
        """Identify the most likely header row within a search window, returning its index, span (1–2 rows), and score."""

        search_limit = min(last_row, first_row + EXCEL_HEADER_SEARCH_ROWS - 1)

        best_row: int | None = None
        best_score = 0.0

        for row_idx in range(first_row, search_limit + 1):
            row_values = [get_value(row_idx, c) for c in range(1, last_col + 1)]
            score = _score_header_row(row_values)

            if score > best_score:
                best_score = score
                best_row = row_idx

        if best_row is None or best_score < EXCEL_HEADER_MIN_SCORE:
            return None, 0, best_score

        span = 1
        next_row = best_row + 1

        if next_row <= last_row:
            header_values = [get_value(best_row, c) for c in range(1, last_col + 1)]
            next_values = [get_value(next_row, c) for c in range(1, last_col + 1)]
            next_score = _score_header_row(next_values)

            next_non_empty_cols = [
                i for i, v in enumerate(next_values) if v is not None and str(v).strip() != ""
            ]

            fills_gaps_only = bool(next_non_empty_cols) and all(
                header_values[i] is None or str(header_values[i]).strip() == ""
                for i in next_non_empty_cols
            )

            if next_score >= EXCEL_HEADER_MIN_SCORE * 0.6 and fills_gaps_only:
                span = 2

        return best_row, span, best_score

    @staticmethod
    def _build_headers(*, get_value, header_rows: list[int], last_col: int,) -> list[str]:
        """Construct column labels by combining values from header rows, falling back to ColX and disambiguating duplicates."""

        headers: list[str] = []
        seen_counts: dict[str, int] = {}

        for col in range(1, last_col + 1):
            parts: list[str] = []

            for row_idx in header_rows:
                value = get_value(row_idx, col)

                if value is not None and str(value).strip() != "":
                    text = _cell_display(value)
                    if text and text not in parts:
                        parts.append(text)

            label = " ".join(parts).strip()

            if not label:
                label = f"Col{get_column_letter(col)}"

            seen_counts[label] = seen_counts.get(label, 0) + 1
            if seen_counts[label] > 1:
                label = f"{label} ({seen_counts[label]})"

            headers.append(label)

        return headers

    def _emit_table_rows(self, *, worksheet: Worksheet, get_value, get_format, headers: list[str],  start_row: int,  last_row: int,  last_col: int,  order: _OrderCounter,  
                         table_index: int,  warnings: list[ParseWarning],  skip_rows: set[int] = frozenset(), is_hidden: bool = False,) -> list[ContentBlock]:
        """Emit table rows as ContentBlock objects, mapping cell values to headers and preserving both display and raw formats."""

        blocks: list[ContentBlock] = []

        end_row = min(last_row, start_row + EXCEL_MAX_DATA_ROWS_PER_SHEET - 1)

        if last_row > end_row:
            warnings.append(
                ParseWarning(
                    code="SHEET_TRUNCATED",
                    message=(
                        f"Sheet '{worksheet.title}' (table #{table_index}): kept "
                        f"{EXCEL_MAX_DATA_ROWS_PER_SHEET} data rows out of "
                        f"{last_row - start_row + 1}; remaining rows were dropped."
                    ),
                    location=SourceLocation(sheet=worksheet.title, table_index=table_index, row=start_row),
                    details={"kept_rows": EXCEL_MAX_DATA_ROWS_PER_SHEET, "total_rows": last_row - start_row + 1},
                )
            )

        col_count = len(headers)
        last_col_letter = get_column_letter(last_col)

        for row_idx in range(start_row, end_row + 1):
            if row_idx in skip_rows:
                continue

            row_values = [get_value(row_idx, c) for c in range(1, last_col + 1)]

            pairs: list[tuple[str, str]] = []
            cell_values: list[str] = []
            cell_values_raw: list[Any] = []

            for col_pos in range(col_count):
                col = col_pos + 1
                raw_value = row_values[col_pos]
                number_format = get_format(row_idx, col)
                display = _cell_display(raw_value, number_format)
                cell_values.append(display)
                cell_values_raw.append(_cell_raw(raw_value))

                header_label = headers[col_pos]

                if display != "" and header_label:
                    pairs.append((header_label, display))

            if not pairs:
                continue  

            raw_text = " | ".join(f"{label}: {value}" for label, value in pairs)
            block_order = order.next()

            blocks.append(
                ContentBlock(
                    block_id=f"excel-{block_order:06d}",
                    block_type=BlockType.TABLE_ROW,
                    raw_text=raw_text,
                    normalized_text=raw_text,
                    order=block_order,
                    location=SourceLocation(
                        sheet=worksheet.title,
                        table_index=table_index,
                        row=row_idx,
                        cell_range=f"A{row_idx}:{last_col_letter}{row_idx}",
                    ),
                    metadata={
                        "headers": headers,
                        "cells": cell_values,
                        "cells_raw": cell_values_raw,
                        "sheet_hidden": is_hidden,
                    },
                )
            )

        return blocks

    def _emit_freeform_rows(self, *, worksheet: Worksheet, get_value, start_row: int, last_row: int, last_col: int, order: _OrderCounter, table_index: int, 
                            warnings: list[ParseWarning], is_hidden: bool = False,) -> list[ContentBlock]:
        """Used when no header row could be confidently detected: emit every non-empty row as-is, without column labels."""

        blocks: list[ContentBlock] = []

        end_row = min(last_row, start_row + EXCEL_MAX_DATA_ROWS_PER_SHEET - 1)

        if last_row > end_row:
            warnings.append(
                ParseWarning(
                    code="SHEET_TRUNCATED",
                    message=(f"Sheet '{worksheet.title}': kept {EXCEL_MAX_DATA_ROWS_PER_SHEET} rows out of {last_row - start_row + 1}; remaining rows were dropped."),
                    location=SourceLocation(sheet=worksheet.title, table_index=table_index, row=start_row),
                    details={"kept_rows": EXCEL_MAX_DATA_ROWS_PER_SHEET, "total_rows": last_row - start_row + 1},
                )
            )

        last_col_letter = get_column_letter(last_col)

        for row_idx in range(start_row, end_row + 1):
            row_values = [get_value(row_idx, c) for c in range(1, last_col + 1)]
            cell_values = [_cell_display(v) for v in row_values]
            cell_values_raw = [_cell_raw(v) for v in row_values]
            non_empty = [v for v in cell_values if v != ""]

            if not non_empty:
                continue

            raw_text = " | ".join(non_empty)
            block_order = order.next()

            blocks.append(
                ContentBlock(
                    block_id=f"excel-{block_order:06d}",
                    block_type=BlockType.TABLE_ROW,
                    raw_text=raw_text,
                    normalized_text=raw_text,
                    order=block_order,
                    location=SourceLocation(
                        sheet=worksheet.title,
                        table_index=table_index,
                        row=row_idx,
                        cell_range=f"A{row_idx}:{last_col_letter}{row_idx}",
                    ),
                    metadata={"cells": cell_values, "cells_raw": cell_values_raw, "sheet_hidden": is_hidden},
                )
            )

        return blocks