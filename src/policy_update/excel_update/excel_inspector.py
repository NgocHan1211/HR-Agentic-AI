"""Read payroll workbook structure without changing the uploaded workbook.

The payroll workbooks used by different customers do not share one fixed
column layout.  In particular, the salary-structure sheets are matrices:

    customer -> contract type -> salary component -> Excel cell

This module turns that layout into explicit, reviewable mapping entries.  A
later Phase 2 service can use an entry to preview an update and, only after
approval, write it to a *new* workbook version.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import re
import unicodedata
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


DEFAULT_MATRIX_SHEETS = ("CƠ CẤU LƯƠNG THÁNG", "CƠ CẤU LƯƠNG NGÀY")


@dataclass(frozen=True)
class SalaryMatrixCell:
    """One configurable amount/formula in a salary-structure matrix."""

    sheet: str
    customer: str
    contract_type: str
    component: str
    cell: str
    value: Any
    formula: str | None = None
    note: str | None = None


@dataclass(frozen=True)
class SheetInspection:
    name: str
    used_range: str
    max_row: int
    max_column: int
    preview_rows: tuple[tuple[Any, ...], ...] = ()
    matrix_cells: tuple[SalaryMatrixCell, ...] = ()
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class WorkbookInspection:
    source_name: str
    sheet_names: tuple[str, ...]
    sheets: tuple[SheetInspection, ...]
    warnings: tuple[str, ...] = field(default_factory=tuple)

    def sheet(self, name: str) -> SheetInspection:
        for item in self.sheets:
            if item.name == name:
                return item
        raise KeyError(f"sheet not found: {name}")


def inspect_workbook(
    file_path: str | Path,
    *,
    matrix_sheets: tuple[str, ...] = DEFAULT_MATRIX_SHEETS,
    preview_rows: int = 8,
) -> WorkbookInspection:
    """Inspect an ``.xlsx``/``.xlsm`` workbook without saving any changes.

    The two known matrix sheets receive richer parsing.  Other sheets still
    return their dimensions and a short preview, which is enough for a future
    mapping UI to let an HR user classify them.
    """

    path = Path(file_path)
    keep_vba = path.suffix.lower() == ".xlsm"
    workbook = load_workbook(path, read_only=False, data_only=False, keep_vba=keep_vba)
    warnings: list[str] = []

    try:
        inspections: list[SheetInspection] = []
        available = set(workbook.sheetnames)
        for name in matrix_sheets:
            if name not in available:
                warnings.append(f"Không tìm thấy sheet cấu trúc lương: {name}")

        for worksheet in workbook.worksheets:
            inspections.append(
                _inspect_sheet(
                    worksheet,
                    is_salary_matrix=worksheet.title in matrix_sheets,
                    preview_limit=preview_rows,
                )
            )
    finally:
        workbook.close()

    return WorkbookInspection(
        source_name=path.name,
        sheet_names=tuple(item.name for item in inspections),
        sheets=tuple(inspections),
        warnings=tuple(warnings),
    )


def find_salary_matrix_cell(
    inspection: WorkbookInspection,
    *,
    sheet: str,
    customer: str,
    contract_type: str,
    component: str,
) -> SalaryMatrixCell:
    """Return one exact matrix cell or raise a clear error for the review UI."""

    target = (_key(customer), _key(contract_type), _key(component))
    matches = [
        item
        for item in inspection.sheet(sheet).matrix_cells
        if (_key(item.customer), _key(item.contract_type), _key(item.component)) == target
    ]
    if not matches:
        raise KeyError(
            "Không tìm thấy mapping Excel cho "
            f"khách hàng={customer!r}, loại hợp đồng={contract_type!r}, khoản={component!r}."
        )
    if len(matches) > 1:
        raise ValueError(
            "Mapping Excel bị trùng cho "
            f"khách hàng={customer!r}, loại hợp đồng={contract_type!r}, khoản={component!r}: "
            + ", ".join(item.cell for item in matches)
        )
    return matches[0]


def _inspect_sheet(
    worksheet: Worksheet,
    *,
    is_salary_matrix: bool,
    preview_limit: int,
) -> SheetInspection:
    max_row, max_column = _used_extent(worksheet)
    used_range = "A1" if not max_row else f"A1:{get_column_letter(max_column)}{max_row}"
    preview = tuple(
        tuple(worksheet.cell(row=row, column=column).value for column in range(1, max_column + 1))
        for row in range(1, min(max_row, preview_limit) + 1)
    )
    warnings: list[str] = []
    matrix_cells: tuple[SalaryMatrixCell, ...] = ()

    if is_salary_matrix:
        matrix_cells, matrix_warnings = _inspect_salary_matrix(worksheet, max_row, max_column)
        warnings.extend(matrix_warnings)

    return SheetInspection(
        name=worksheet.title,
        used_range=used_range,
        max_row=max_row,
        max_column=max_column,
        preview_rows=preview,
        matrix_cells=matrix_cells,
        warnings=tuple(warnings),
    )


def _inspect_salary_matrix(
    worksheet: Worksheet,
    max_row: int,
    max_column: int,
) -> tuple[tuple[SalaryMatrixCell, ...], list[str]]:
    """Parse a customer/contract header salary matrix.

    Some workbooks begin at A1, while others keep one or more title/blank
    rows and columns.  Locate the contract header from its ``NOTE`` column,
    then derive the customer row and component column from that location.
    """

    warnings: list[str] = []
    if max_row < 3 or max_column < 2:
        return (), [f"Sheet '{worksheet.title}' quá nhỏ để nhận diện bảng cơ cấu lương."]

    merged_values = _merged_values(worksheet)
    contract_header_row = _find_contract_header_row(worksheet, max_row, max_column, merged_values)
    if contract_header_row is None:
        return (), [
            f"Không tìm thấy hàng loại hợp đồng trong '{worksheet.title}'. "
            "Cần có ít nhất một cột NOTE hoặc nhãn Thời vụ/Chính thức."
        ]
    customer_header_row = contract_header_row - 1
    contract_columns = [
        column
        for column in range(1, max_column + 1)
        if _text(_value(worksheet, contract_header_row, column, merged_values))
        and _key(_text(_value(worksheet, contract_header_row, column, merged_values))) != "note"
    ]
    if not contract_columns:
        return (), [f"Không tìm thấy cột giá trị hợp đồng trong '{worksheet.title}'."]
    component_column = min(contract_columns) - 1
    if component_column < 1:
        return (), [f"Không xác định được cột tên khoản lương trong '{worksheet.title}'."]

    entries: list[SalaryMatrixCell] = []
    current_customer = ""

    for column in range(component_column + 1, max_column + 1):
        customer = _text(_value(worksheet, customer_header_row, column, merged_values))
        if customer:
            current_customer = customer
        else:
            customer = current_customer

        contract_type = _text(_value(worksheet, contract_header_row, column, merged_values))
        if not customer or not contract_type:
            continue

        # NOTE columns describe a rule but are not safe targets for a numeric
        # update.  They are attached to the two preceding value columns below.
        if _key(contract_type) == "note":
            continue

        note = _nearest_note(worksheet, contract_header_row, column, merged_values)
        for row in range(contract_header_row + 1, max_row + 1):
            component = _text(_value(worksheet, row, component_column, merged_values))
            if not component:
                continue

            cell = worksheet.cell(row=row, column=column)
            value = cell.value
            if value is None or (isinstance(value, str) and not value.strip()):
                continue

            formula = value if isinstance(value, str) and value.startswith("=") else None
            entries.append(
                SalaryMatrixCell(
                    sheet=worksheet.title,
                    customer=customer,
                    contract_type=contract_type,
                    component=component,
                    cell=cell.coordinate,
                    value=value,
                    formula=formula,
                    note=note,
                )
            )

    if not entries:
        warnings.append(
            f"Không đọc được giá trị từ ma trận '{worksheet.title}'. "
            "Kiểm tra lại hàng khách hàng, hàng loại hợp đồng và cột khoản lương."
        )
    return tuple(entries), warnings


def _find_contract_header_row(
    worksheet: Worksheet,
    max_row: int,
    max_column: int,
    merged_values: dict[tuple[int, int], Any],
) -> int | None:
    """Find the row naming contract types, such as ``Chính thức`` and NOTE."""

    search_limit = min(max_row, 12)
    best: tuple[int, int] | None = None
    for row in range(1, search_limit + 1):
        values = [_key(_text(_value(worksheet, row, column, merged_values))) for column in range(1, max_column + 1)]
        score = sum(value == "note" for value in values) * 10
        score += sum("thoi vu" in value or "chinh thuc" in value for value in values)
        if score and (best is None or score > best[1]):
            best = (row, score)
    return best[0] if best else None


def _nearest_note(
    worksheet: Worksheet,
    header_row: int,
    column: int,
    merged_values: dict[tuple[int, int], Any],
) -> str | None:
    """Use the nearest NOTE column in the same customer group, when present."""

    for candidate in (column + 1, column + 2):
        if candidate > worksheet.max_column:
            continue
        if _key(_value(worksheet, header_row, candidate, merged_values)) == "note":
            return get_column_letter(candidate)
    return None


def _used_extent(worksheet: Worksheet) -> tuple[int, int]:
    last_row = 0
    last_column = 0
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip():
                last_row = max(last_row, cell.row)
                last_column = max(last_column, cell.column)
    return last_row, last_column


def _merged_values(worksheet: Worksheet) -> dict[tuple[int, int], Any]:
    values: dict[tuple[int, int], Any] = {}
    for merged_range in worksheet.merged_cells.ranges:
        value = worksheet.cell(merged_range.min_row, merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for column in range(merged_range.min_col, merged_range.max_col + 1):
                values[(row, column)] = value
    return values


def _value(worksheet: Worksheet, row: int, column: int, merged_values: dict[tuple[int, int], Any]) -> Any:
    value = worksheet.cell(row=row, column=column).value
    return value if value is not None else merged_values.get((row, column))


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _key(value: str) -> str:
    plain = "".join(
        char for char in unicodedata.normalize("NFD", value.casefold()) if unicodedata.category(char) != "Mn"
    )
    return re.sub(r"[^a-z0-9]+", " ", plain).strip()
