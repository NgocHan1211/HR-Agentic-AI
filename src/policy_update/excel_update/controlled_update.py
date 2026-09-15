"""Controlled Phase-2 updates for payroll configuration workbooks.

This module deliberately separates *proposal* from *mutation*: an AI/ChangeSet
may propose a new allowance, but only an APPROVED ChangeSet with an exact,
reviewable Excel target can be applied.  The source workbook is never edited.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
from shutil import copy2
from typing import Any, Iterable
from uuid import uuid4

from openpyxl import load_workbook

from ..change_management.changeset_schema import (
    ChangeCategory,
    ChangeSet,
    ChangeSetStatus,
    UpdateOperation,
)
from ..change_management.state_machine import Actor, transition
from .excel_inspector import WorkbookInspection, find_salary_matrix_cell


@dataclass(frozen=True)
class ExcelTargetSelector:
    """Human-reviewable selector for a cell in a salary configuration matrix."""

    sheet: str
    customer: str
    contract_type: str
    component: str


@dataclass(frozen=True)
class PreviewItem:
    operation_id: str
    sheet: str | None
    cell: str | None
    before: Any
    after: Any
    status: str  # READY | CONFLICT | INVALID_TARGET
    message: str = ""


@dataclass(frozen=True)
class UpdatePreview:
    source_path: str
    changeset_id: str
    items: tuple[PreviewItem, ...]

    @property
    def is_safe_to_apply(self) -> bool:
        return bool(self.items) and all(item.status == "READY" for item in self.items)


@dataclass
class WorkbookUpdateRun:
    """Audit record for one non-destructive update run."""

    changeset_id: str
    idempotency_key: str
    source_path: str
    backup_path: str
    output_path: str | None
    source_sha256: str
    output_sha256: str | None
    operation_ids: tuple[str, ...]
    status: str  # COMPLETED | FAILED | ROLLED_BACK
    id: str = field(default_factory=lambda: f"update-run-{uuid4().hex[:12]}")
    created_at: datetime = field(default_factory=lambda: datetime.now(timezone.utc))
    error: str | None = None
    rollback_path: str | None = None


@dataclass
class UpdateRunRegistry:
    """Small in-memory idempotency registry; replace with the project DB in deployment."""

    by_idempotency_key: dict[str, WorkbookUpdateRun] = field(default_factory=dict)

    def get(self, key: str) -> WorkbookUpdateRun | None:
        return self.by_idempotency_key.get(key)

    def save(self, run: WorkbookUpdateRun) -> None:
        self.by_idempotency_key[run.idempotency_key] = run


def resolve_parameter_change(
    inspection: WorkbookInspection,
    change_item: Any,
    selector: ExcelTargetSelector,
    *,
    workbook_id: str | None = None,
) -> UpdateOperation:
    """Resolve a reviewed parameter change to one exact Excel cell.

    Formula changes remain owned by FormulaSpec review.  This resolver is for
    a policy parameter such as an allowance/rate amount, where the desired new
    value is explicit and must match the workbook's current cell before apply.
    """
    if change_item.category is not ChangeCategory.PARAMETER_CHANGE:
        raise ValueError("only PARAMETER_CHANGE may be resolved to a direct Excel cell")
    if not change_item.is_ready_for_update_operation:
        raise ValueError("ChangeItem is missing evidence, effective date, field_path, or is blocked")
    target = find_salary_matrix_cell(
        inspection,
        sheet=selector.sheet,
        customer=selector.customer,
        contract_type=selector.contract_type,
        component=selector.component,
    )
    return UpdateOperation(
        change_item_id=change_item.id,
        target_type="cell",
        workbook_id=workbook_id or inspection.source_name,
        sheet=target.sheet,
        cell_or_row_key=target.cell,
        before=target.value,
        after=change_item.proposed_value,
        data_type=_data_type(change_item.proposed_value),
        precondition={"expected_value": target.value, "selector": selector.__dict__},
    )


def preview_excel_update(source_path: str | Path, changeset: ChangeSet, operations: Iterable[UpdateOperation]) -> UpdatePreview:
    """Dry-run every operation against the actual source file; never writes it."""
    path = Path(source_path)
    items: list[PreviewItem] = []
    workbook = load_workbook(path, read_only=True, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    try:
        for operation in operations:
            items.append(_preview_item(workbook, operation))
    finally:
        workbook.close()
    return UpdatePreview(source_path=str(path), changeset_id=changeset.id, items=tuple(items))


def apply_approved_changeset(
    source_path: str | Path,
    changeset: ChangeSet,
    operations: Iterable[UpdateOperation],
    output_directory: str | Path,
    *,
    registry: UpdateRunRegistry | None = None,
) -> WorkbookUpdateRun:
    """Apply approved, conflict-free cell updates to a new workbook version.

    A backup and output are copied before mutation.  The source file remains
    untouched.  The output is re-opened and reconciled before completion.
    """
    if changeset.status is not ChangeSetStatus.APPROVED:
        raise ValueError("only an APPROVED ChangeSet may be applied")
    if changeset.blocks_release:
        raise ValueError("ChangeSet contains a release-blocking item")
    operations = tuple(operations)
    _validate_operations(changeset, operations)
    registry = registry or UpdateRunRegistry()
    existing = registry.get(changeset.idempotency_key)
    if existing is not None:
        return existing

    source = Path(source_path)
    if not source.is_file():
        raise FileNotFoundError(source)
    output_dir = Path(output_directory)
    output_dir.mkdir(parents=True, exist_ok=True)
    suffix = source.suffix.lower()
    token = f"{changeset.id}-{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}"
    backup = output_dir / f"{source.stem}-{token}-backup{suffix}"
    output = output_dir / f"{source.stem}-{token}-updated{suffix}"
    copy2(source, backup)
    copy2(source, output)
    source_hash = _sha256(source)

    preview = preview_excel_update(source, changeset, operations)
    if not preview.is_safe_to_apply:
        message = "; ".join(f"{item.operation_id}: {item.message}" for item in preview.items if item.status != "READY")
        run = WorkbookUpdateRun(
            changeset_id=changeset.id, idempotency_key=changeset.idempotency_key,
            source_path=str(source), backup_path=str(backup), output_path=None,
            source_sha256=source_hash, output_sha256=None,
            operation_ids=tuple(item.id for item in operations), status="FAILED", error=message,
        )
        registry.save(run)
        raise ValueError(f"preview failed; no workbook was updated: {message}")

    transition(changeset, ChangeSetStatus.APPLYING, Actor.SYSTEM)
    try:
        workbook = load_workbook(output, read_only=False, data_only=False, keep_vba=suffix == ".xlsm")
        try:
            for operation in operations:
                worksheet = workbook[operation.sheet or ""]
                worksheet[operation.cell_or_row_key or ""].value = operation.after
            workbook.save(output)
        finally:
            workbook.close()
        _reconcile(output, operations)
        transition(changeset, ChangeSetStatus.COMPLETED, Actor.SYSTEM)
        run = WorkbookUpdateRun(
            changeset_id=changeset.id, idempotency_key=changeset.idempotency_key,
            source_path=str(source), backup_path=str(backup), output_path=str(output),
            source_sha256=source_hash, output_sha256=_sha256(output),
            operation_ids=tuple(item.id for item in operations), status="COMPLETED",
        )
    except Exception as exc:
        if changeset.status is ChangeSetStatus.APPLYING:
            transition(changeset, ChangeSetStatus.FAILED, Actor.SYSTEM)
        run = WorkbookUpdateRun(
            changeset_id=changeset.id, idempotency_key=changeset.idempotency_key,
            source_path=str(source), backup_path=str(backup), output_path=None,
            source_sha256=source_hash, output_sha256=None,
            operation_ids=tuple(item.id for item in operations), status="FAILED", error=str(exc),
        )
    registry.save(run)
    return run


def rollback_update_run(
    run: WorkbookUpdateRun,
    changeset: ChangeSet,
    output_directory: str | Path,
    *,
    actor: Actor = Actor.APPROVER,
) -> WorkbookUpdateRun:
    """Create a restored copy from the immutable backup; never overwrite source/output."""
    if run.status not in {"COMPLETED", "FAILED"}:
        raise ValueError(f"cannot rollback run in status {run.status}")
    backup = Path(run.backup_path)
    if not backup.is_file():
        raise FileNotFoundError(f"backup is missing: {backup}")
    destination_dir = Path(output_directory)
    destination_dir.mkdir(parents=True, exist_ok=True)
    restored = destination_dir / f"{backup.stem}-restored{backup.suffix}"
    copy2(backup, restored)
    if changeset.status in {ChangeSetStatus.COMPLETED, ChangeSetStatus.FAILED}:
        transition(changeset, ChangeSetStatus.ROLLED_BACK, actor)
    run.status = "ROLLED_BACK"
    run.rollback_path = str(restored)
    return run


def _preview_item(workbook: Any, operation: UpdateOperation) -> PreviewItem:
    if operation.target_type != "cell" or not operation.sheet or not operation.cell_or_row_key:
        return PreviewItem(operation.id, operation.sheet, operation.cell_or_row_key, operation.before, operation.after,
                           "INVALID_TARGET", "only exact cell targets are supported")
    if operation.sheet not in workbook.sheetnames:
        return PreviewItem(operation.id, operation.sheet, operation.cell_or_row_key, operation.before, operation.after,
                           "INVALID_TARGET", f"sheet {operation.sheet!r} does not exist")
    actual = workbook[operation.sheet][operation.cell_or_row_key].value
    expected = operation.precondition.get("expected_value", operation.before)
    if not _same_value(actual, expected):
        return PreviewItem(operation.id, operation.sheet, operation.cell_or_row_key, actual, operation.after,
                           "CONFLICT", f"expected {expected!r}, found {actual!r}; workbook baseline changed")
    return PreviewItem(operation.id, operation.sheet, operation.cell_or_row_key, actual, operation.after, "READY")


def _validate_operations(changeset: ChangeSet, operations: tuple[UpdateOperation, ...]) -> None:
    if not operations:
        raise ValueError("at least one UpdateOperation is required")
    seen_targets: set[tuple[str | None, str | None]] = set()
    item_ids = {item.id for item in changeset.items}
    for operation in operations:
        if operation.change_item_id not in item_ids:
            raise ValueError(f"operation {operation.id} does not belong to ChangeSet {changeset.id}")
        target = (operation.sheet, operation.cell_or_row_key)
        if target in seen_targets:
            raise ValueError(f"multiple operations target the same Excel cell: {target}")
        seen_targets.add(target)


def _reconcile(output_path: Path, operations: tuple[UpdateOperation, ...]) -> None:
    workbook = load_workbook(output_path, read_only=True, data_only=False, keep_vba=output_path.suffix.lower() == ".xlsm")
    try:
        for operation in operations:
            actual = workbook[operation.sheet or ""][operation.cell_or_row_key or ""].value
            if not _same_value(actual, operation.after):
                raise ValueError(
                    f"reconciliation failed for {operation.sheet}!{operation.cell_or_row_key}: "
                    f"expected {operation.after!r}, found {actual!r}"
                )
    finally:
        workbook.close()


def _same_value(left: Any, right: Any) -> bool:
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return abs(float(left) - float(right)) < 1e-9
    return left == right


def _data_type(value: Any) -> str:
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    return "text"


def _sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()
