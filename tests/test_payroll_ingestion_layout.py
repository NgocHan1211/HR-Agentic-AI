from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from payroll.ingestion import SheetMappingSpec, normalize_attendance, parse_attendance_excel


def test_attendance_ingestion_supports_multiline_headers_and_configured_row_selection(tmp_path: Path) -> None:
    path = tmp_path / "grouped_attendance.xlsx"
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Payroll"
    sheet.cell(10, 1, "Mã NV"); sheet.cell(10, 2, "Lọc"); sheet.cell(10, 3, "Ngày công")
    sheet.cell(11, 3, "thực tế")
    sheet.cell(12, 1, "dòng kỹ thuật")
    sheet.cell(13, 1, "mã cột kỹ thuật")
    sheet.append(["E-01", "CÔNG", 22])
    sheet.append(["E-01", "NGHỈ", 0])
    sheet.append(["E-01", "TĂNG CA", 0])
    workbook.save(path)

    spec = SheetMappingSpec("C1", "attendance", {"Payroll": {
        "header_rows": [10, 11], "data_start_row": 14,
        "row_selector": {"column": "Lọc", "equals": "CÔNG"},
        "columns": {"Mã NV": "employee_id", "Ngày công | thực tế": "days_worked_in_month"},
    }})
    raw = parse_attendance_excel(path, spec)

    assert list(raw["Payroll"].columns[:3]) == ["Mã NV", "Lọc", "Ngày công | thực tế"]
    assert raw["Payroll"].to_dict("records") == [{"Mã NV": "E-01", "Lọc": "CÔNG", "Ngày công | thực tế": 22}]
    records = normalize_attendance(raw, spec, "2026-07")
    assert records[0].to_dict()["days_worked_in_month"] == 22.0


def test_row_selector_is_template_configuration_not_a_hard_coded_rule(tmp_path: Path) -> None:
    path = tmp_path / "alternative_roles.xlsx"
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Attendance"
    sheet.append(["Employee", "Role", "Hours"])
    sheet.append(["E-01", "PRIMARY", 8])
    sheet.append(["E-01", "DETAIL", 0])
    workbook.save(path)

    spec = SheetMappingSpec("C1", "attendance", {"Attendance": {
        "row_selector": {"column": "Role", "equals": "PRIMARY"},
    }})
    raw = parse_attendance_excel(path, spec)
    assert raw["Attendance"].to_dict("records") == [{"Employee": "E-01", "Role": "PRIMARY", "Hours": 8}]
