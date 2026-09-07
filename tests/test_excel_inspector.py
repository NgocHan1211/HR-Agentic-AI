from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook

from policy_update.excel_update import find_salary_matrix_cell, inspect_workbook


class ExcelInspectorTests(unittest.TestCase):
    def test_reads_customer_contract_component_matrix(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "salary.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "CƠ CẤU LƯƠNG THÁNG"
            worksheet.merge_cells("B1:D1")
            worksheet["B1"] = "ACME"
            worksheet["B2"] = "Thời vụ"
            worksheet["C2"] = "Chính thức"
            worksheet["D2"] = "NOTE"
            worksheet["A3"] = "LƯƠNG CƠ BẢN"
            worksheet["B3"] = 4_000_000
            worksheet["C3"] = 5_000_000
            worksheet["D3"] = "Tính theo ngày công"
            worksheet["A4"] = "PHỤ CẤP CHUYÊN CẦN"
            worksheet["B4"] = 300_000
            worksheet["C4"] = 500_000
            workbook.save(path)

            inspection = inspect_workbook(path)
            entry = find_salary_matrix_cell(
                inspection,
                sheet="CƠ CẤU LƯƠNG THÁNG",
                customer="acme",
                contract_type="chinh thuc",
                component="phu cap chuyen can",
            )

            self.assertEqual("C4", entry.cell)
            self.assertEqual(500_000, entry.value)
            self.assertEqual("D", entry.note)

    def test_keeps_excel_formula_as_a_reviewable_value(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "salary.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "CƠ CẤU LƯƠNG THÁNG"
            worksheet["B1"] = "ACME"
            worksheet["B2"] = "Chính thức"
            worksheet["A3"] = "BHXH 10.5%"
            worksheet["B3"] = "=5000000*10.5%"
            workbook.save(path)

            entry = find_salary_matrix_cell(
                inspect_workbook(path),
                sheet="CƠ CẤU LƯƠNG THÁNG",
                customer="ACME",
                contract_type="Chính thức",
                component="BHXH 10.5%",
            )

            self.assertEqual("=5000000*10.5%", entry.formula)


if __name__ == "__main__":
    unittest.main()
