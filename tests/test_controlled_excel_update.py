from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from policy_update.change_management.changeset_schema import (
    ChangeCategory,
    ChangeItem,
    ChangeSet,
    ChangeSetStatus,
)
from policy_update.excel_update import (
    ExcelTargetSelector,
    apply_approved_changeset,
    inspect_workbook,
    preview_excel_update,
    resolve_parameter_change,
    rollback_update_run,
    simulate_changeset_impact,
)


class ControlledExcelUpdateTests(unittest.TestCase):
    def _workbook(self, directory: str) -> Path:
        path = Path(directory) / "master.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "CƠ CẤU LƯƠNG THÁNG"
        sheet.merge_cells("B1:D1")
        sheet["B1"] = "ACME"
        sheet["B2"] = "Thời vụ"
        sheet["C2"] = "Chính thức"
        sheet["D2"] = "NOTE"
        sheet["A3"] = "PHỤ CẤP CHUYÊN CẦN"
        sheet["B3"] = 300_000
        sheet["C3"] = 450_000
        sheet["D3"] = "Theo hợp đồng"
        book.save(path)
        return path

    def _approved_changeset(self) -> tuple[ChangeSet, ChangeItem]:
        changeset = ChangeSet(
            company_id="ACME", baseline_formula_version=1,
            baseline_workbook_version="v1", effective_from=date(2026, 9, 1), scope="Chính thức",
        )
        item = ChangeItem(
            changeset_id=changeset.id, category=ChangeCategory.PARAMETER_CHANGE,
            field_path="Allowances[code=ATTENDANCE].amount", old_value=450_000, proposed_value=500_000,
            reason="Policy allowance change", evidence_refs=("evidence-1",), effective_date=date(2026, 9, 1),
            confidence=0.95,
        )
        changeset.add_item(item)
        changeset.status = ChangeSetStatus.APPROVED
        return changeset, item

    def test_preview_apply_reconcile_and_rollback_without_touching_source(self) -> None:
        with TemporaryDirectory() as directory:
            source = self._workbook(directory)
            changeset, item = self._approved_changeset()
            operation = resolve_parameter_change(
                inspect_workbook(source), item,
                ExcelTargetSelector("CƠ CẤU LƯƠNG THÁNG", "ACME", "Chính thức", "PHỤ CẤP CHUYÊN CẦN"),
            )
            preview = preview_excel_update(source, changeset, [operation])
            self.assertTrue(preview.is_safe_to_apply)
            run = apply_approved_changeset(source, changeset, [operation], Path(directory) / "runs")
            self.assertEqual("COMPLETED", run.status)
            self.assertEqual(ChangeSetStatus.COMPLETED, changeset.status)
            self.assertEqual(450_000, load_workbook(source, data_only=False)["CƠ CẤU LƯƠNG THÁNG"]["C3"].value)
            self.assertEqual(500_000, load_workbook(run.output_path, data_only=False)["CƠ CẤU LƯƠNG THÁNG"]["C3"].value)
            rolled_back = rollback_update_run(run, changeset, Path(directory) / "runs")
            self.assertEqual("ROLLED_BACK", rolled_back.status)
            self.assertEqual(450_000, load_workbook(rolled_back.rollback_path, data_only=False)["CƠ CẤU LƯƠNG THÁNG"]["C3"].value)

    def test_preview_detects_baseline_conflict(self) -> None:
        with TemporaryDirectory() as directory:
            source = self._workbook(directory)
            changeset, item = self._approved_changeset()
            operation = resolve_parameter_change(
                inspect_workbook(source), item,
                ExcelTargetSelector("CƠ CẤU LƯƠNG THÁNG", "ACME", "Chính thức", "PHỤ CẤP CHUYÊN CẦN"),
            )
            book = load_workbook(source)
            book["CƠ CẤU LƯƠNG THÁNG"]["C3"] = 475_000
            book.save(source)
            preview = preview_excel_update(source, changeset, [operation])
            self.assertFalse(preview.is_safe_to_apply)
            self.assertEqual("CONFLICT", preview.items[0].status)

    def test_impact_report_reconciles_employee_results(self) -> None:
        report = simulate_changeset_impact(
            [{"employee_id": "NV001", "gross_salary": 10_000_000, "net_salary": 9_000_000}],
            [{"employee_id": "NV001", "gross_salary": 10_500_000, "net_salary": 9_450_000}],
        )
        self.assertEqual(1, report.affected_employee_count)
        self.assertEqual(500_000, report.total_gross_delta)
        self.assertEqual(450_000, report.total_net_delta)


if __name__ == "__main__":
    unittest.main()
